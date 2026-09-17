import json
import math
import numpy as np
import openpyxl
import pandas as pd
import pytest
from pathlib import Path

from app.analytics.trends import TrendAnalyzer
from app.anomaly.detector import AnomalyDetector, AnomalyEngine
from app.database.safety import SQLSafetyValidator, SQLSafetyError
from app.metrics.kpi_engine import kpi_engine
from app.quality.engine import DataQualityEngine
from app.reports.builder import ReportBuilder
from app.reports.data_exporter import DataExporter, _sanitize_for_json
from app.reports.excel_exporter import ExcelExporter
from app.reports.pdf_exporter import PDFExporter, _ensure_unicode_fonts
from app.rules.rule_engine import business_rule_engine


def test_sql_identifier_validation_blocks_injections():
    malicious_identifiers = [
        'users"; DROP TABLE users; --',
        'orders" UNION SELECT password FROM admin_users --',
        "table name with spaces",
        "table;name",
        "table-name",
        "123table",
        "table'name",
        'table"name',
        "`table`",
    ]
    for ident in malicious_identifiers:
        with pytest.raises(SQLSafetyError):
            SQLSafetyValidator.validate_identifier(ident)

    # Valid identifiers
    assert SQLSafetyValidator.validate_identifier("orders") == "orders"
    assert SQLSafetyValidator.validate_identifier("user_accounts_2026") == "user_accounts_2026"
    assert SQLSafetyValidator.validate_identifier("_internal_id") == "_internal_id"


def test_sql_table_identifier_supports_schemas_and_blocks_injections():
    assert SQLSafetyValidator.validate_table_identifier("orders") == "orders"
    assert SQLSafetyValidator.validate_table_identifier("public.orders") == "public.orders"
    assert SQLSafetyValidator.validate_table_identifier("main.customers") == "main.customers"

    malicious_tables = [
        'public."orders"; DROP TABLE orders;',
        "public.orders.extra",
        "orders; DROP TABLE orders;",
        'orders" UNION SELECT 1 --',
    ]
    for tbl in malicious_tables:
        with pytest.raises(SQLSafetyError):
            SQLSafetyValidator.validate_table_identifier(tbl)


def test_kpi_formula_validation():
    # Valid formulas
    valid_formulas = [
        "SUM(total_amount)",
        "COUNT(id)",
        "COUNT(*)",
        "SUM(total_amount) / COUNT(id)",
        "AVG(defect_rate)",
        "MAX(operating_temp_c) - MIN(operating_temp_c)",
        "SUM(revenue) * 1.18",
    ]
    for formula in valid_formulas:
        is_safe, err, cols = SQLSafetyValidator.validate_kpi_formula(formula)
        assert is_safe, f"Expected formula '{formula}' to be valid, error: {err}"

    # Verify column extraction
    _, _, cols = SQLSafetyValidator.validate_kpi_formula("SUM(total_amount) / COUNT(id)")
    assert "total_amount" in cols
    assert "id" in cols
    assert "SUM" not in cols
    assert "COUNT" not in cols

    # Malicious injection formulas
    malicious_formulas = [
        "1; DROP TABLE orders --",
        "SUM(total_amount); DELETE FROM orders",
        "SUM(total_amount) UNION SELECT password FROM users",
        "EXEC xp_cmdshell('dir')",
        "SUM(amount) -- sql comment",
        "SUM(amount) /* block comment */",
        "SUM(amount) 'literal string'",
    ]
    for formula in malicious_formulas:
        is_safe, err, _ = SQLSafetyValidator.validate_kpi_formula(formula)
        assert not is_safe, f"Expected malicious formula '{formula}' to be blocked!"


def test_condition_sql_validation():
    valid_conditions = [
        "defect_rate > 0.04",
        "stock_level < 10 AND status = 'ACTIVE'",
        "quantity_on_hand <= minimum_stock",
        "salary >= 50000 OR department = 'Sales'",
    ]
    for cond in valid_conditions:
        is_safe, err = SQLSafetyValidator.validate_condition_sql(cond)
        assert is_safe, f"Expected condition '{cond}' to be safe, error: {err}"

    malicious_conditions = [
        "1=1; DROP TABLE orders;",
        "defect_rate > 0.04; DELETE FROM production",
        "1=1 -- bypass where",
        "1=1 /* comment */",
    ]
    for cond in malicious_conditions:
        is_safe, err = SQLSafetyValidator.validate_condition_sql(cond)
        assert not is_safe, f"Expected condition '{cond}' to be blocked!"


def test_kpi_engine_rejects_malicious_formula():
    with pytest.raises(SQLSafetyError):
        kpi_engine.add_kpi(
            name="Injected KPI",
            database_name="demo_enterprise",
            table_name="orders",
            formula="1; DROP TABLE orders --",
        )


def test_rule_engine_rejects_malicious_rule():
    with pytest.raises(SQLSafetyError):
        business_rule_engine.create_rule(
            name="Injected Rule",
            database_name="demo_enterprise",
            table_name="orders",
            condition_sql="1=1; DROP TABLE orders;",
        )


def test_anomaly_detection_inf_and_nan_resilience():
    # Test series with np.inf, -np.inf, and np.nan mixed with normal values
    data = [10.0, 12.0, 11.0, 9.0, 10.5, 11.2, 10.8, 10.1, 9.8, 10.3, 100.0, np.inf, -np.inf, np.nan]
    s = pd.Series(data)

    # Z-Score
    z_anomalies = AnomalyDetector.detect_zscore(s, metric_name="test_metric")
    assert len(z_anomalies) > 0
    top_z = z_anomalies[0]
    assert top_z.observed_value == 100.0
    assert not math.isnan(top_z.score)
    assert not math.isinf(top_z.score)

    # IQR
    iqr_anomalies = AnomalyDetector.detect_iqr(s, metric_name="test_metric")
    assert len(iqr_anomalies) > 0
    top_iqr = iqr_anomalies[0]
    assert top_iqr.observed_value == 100.0
    assert not math.isnan(top_iqr.score)
    assert not math.isinf(top_iqr.score)


def test_isolation_forest_adaptive_contamination():
    # Small dataset with 12 samples (previously could cause issues with contamination=0.03 where 12*0.03 < 1)
    df = pd.DataFrame({
        "metric1": [10.0, 11.0, 10.5, 10.2, 10.8, 11.1, 10.3, 10.7, 10.4, 10.9, 10.6, 95.0],
        "metric2": [20.0, 21.0, 20.5, 20.2, 20.8, 21.1, 20.3, 20.7, 20.4, 20.9, 20.6, 195.0],
    })
    anomalies = AnomalyDetector.detect_isolation_forest(df, numeric_cols=["metric1", "metric2"], contamination=0.03)
    # Should complete without error and identify the outlier
    assert len(anomalies) > 0
    assert anomalies[0].observed_value == 95.0


def test_rolling_window_duplicate_index_resilience():
    # Non-unique index
    df = pd.DataFrame({
        "time": pd.date_range("2026-01-01", periods=10, freq="D"),
        "val": [10.0, 10.2, 10.1, 10.3, 10.2, 10.4, 10.1, 10.5, 80.0, 10.2],
    }, index=[0, 0, 1, 1, 2, 2, 3, 3, 4, 4])

    anomalies = AnomalyDetector.detect_rolling(df, metric_col="val", time_col="time", window=4)
    assert len(anomalies) > 0
    assert anomalies[0].observed_value == 80.0


def test_trend_p_value_strict_significance():
    # Random noisy series without significant momentum (slope flat, p > 0.05)
    noise_series = [
        {"period": "2026-01", "value": 100.0},
        {"period": "2026-02", "value": 102.0},
        {"period": "2026-03", "value": 99.0},
        {"period": "2026-04", "value": 101.5},
        {"period": "2026-05", "value": 100.5},
    ]
    res = TrendAnalyzer.analyze_series(noise_series, metric_name="Noise")
    assert res["trend_direction"] == "STABLE"


def test_pdf_unicode_turkish_characters():
    reg, bold = _ensure_unicode_fonts()
    assert reg == "DejaVuSans"
    assert bold == "DejaVuSans-Bold"

    builder = ReportBuilder("demo_enterprise")
    data = builder.build_report_data(
        title="Türkçe Karakter Test Raporu (Ğ, Ü, Ş, İ, Ö, Ç, ğ, ü, ş, ı, ö, ç)",
        period="Eylül 2026",
        language="tr",
    )
    pdf_path = PDFExporter.export(data)
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 1000


def test_excel_numeric_percentage_formatting():
    builder = ReportBuilder("demo_enterprise")
    data = builder.build_report_data(title="Excel Numeric Test", period="Q3 2026", language="en")
    xlsx_path = ExcelExporter.export(data)
    assert xlsx_path.exists()

    wb = openpyxl.load_workbook(xlsx_path)
    ws_kpi = wb["KPIs"]

    # Row 2, column 6 is MoM growth
    cell_mom = ws_kpi.cell(row=2, column=6)
    if cell_mom.value != "-":
        assert isinstance(cell_mom.value, (int, float)), f"Expected float, got {type(cell_mom.value)}"
        assert "%" in cell_mom.number_format

    # Row 2, column 8 is target achievement
    cell_ach = ws_kpi.cell(row=2, column=8)
    if cell_ach.value != "-":
        assert isinstance(cell_ach.value, (int, float)), f"Expected float, got {type(cell_ach.value)}"
        assert "%" in cell_ach.number_format

    # Trends sheet
    ws_trends = wb["Trends"]
    cell_tg = ws_trends.cell(row=2, column=3)
    if cell_tg.value != "-":
        assert isinstance(cell_tg.value, (int, float))
        assert "%" in cell_tg.number_format


def test_json_nan_inf_sanitization():
    raw_data = {
        "normal_val": 42.5,
        "nan_val": float("nan"),
        "inf_val": float("inf"),
        "neg_inf_val": float("-inf"),
        "nested": {
            "sub_nan": float("nan"),
            "sub_list": [1.0, float("nan"), float("inf")],
        },
    }
    sanitized = _sanitize_for_json(raw_data)
    assert sanitized["nan_val"] is None
    assert sanitized["inf_val"] is None
    assert sanitized["neg_inf_val"] is None
    assert sanitized["nested"]["sub_nan"] is None
    assert sanitized["nested"]["sub_list"] == [1.0, None, None]

    # Standard json dumps without allow_nan should now succeed
    json_str = json.dumps(sanitized)
    assert "NaN" not in json_str
    assert "Infinity" not in json_str
    parsed = json.loads(json_str)
    assert parsed["nan_val"] is None


def test_quality_engine_turkish_metadata():
    engine = DataQualityEngine("demo_enterprise")
    res = engine.evaluate_table("orders", sample_size=500)
    assert "quality_grade_tr" in res
    assert res["quality_grade_tr"] in ["MÜKEMMEL", "İYİ", "ORTA", "ZAYIF"]
    assert "scoring_explanation_tr" in res
    assert "temel puan" in res["scoring_explanation_tr"]
