from pathlib import Path
from typing import Any, Dict, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.core.config import settings


class ExcelExporter:
    """Generates multi-worksheet corporate Excel report using openpyxl."""

    HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="0F172A")
    BOLD_FONT = Font(name="Calibri", size=11, bold=True)
    REGULAR_FONT = Font(name="Calibri", size=11)
    THIN_BORDER = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    @classmethod
    def export(cls, report_data: Dict[str, Any], output_path: Optional[Path] = None) -> Path:
        dest = output_path or (
            settings.UDI_REPORTS_DIR / "exports" / f"report_{report_data.get('database_name', 'db')}_{int(hash(report_data.get('generated_at', '')) % 1000000)}.xlsx"
        )
        dest.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # 1. Summary Sheet
        ws_summary = wb.create_sheet(title="Summary")
        cls._populate_summary(ws_summary, report_data)

        # 2. KPIs Sheet
        ws_kpis = wb.create_sheet(title="KPIs")
        cls._populate_kpis(ws_kpis, report_data.get("kpis", []))

        # 3. Trends Sheet
        ws_trends = wb.create_sheet(title="Trends")
        cls._populate_trends(ws_trends, report_data.get("trends", []))

        # 4. Anomalies Sheet
        ws_anom = wb.create_sheet(title="Anomalies")
        cls._populate_anomalies(ws_anom, report_data.get("anomalies", []))

        # 5. Data Quality Sheet
        ws_dq = wb.create_sheet(title="Data Quality")
        cls._populate_data_quality(ws_dq, report_data.get("quality", {}))

        # 6. Raw Results / Violations Sheet
        ws_raw = wb.create_sheet(title="Raw Results")
        cls._populate_raw_results(ws_raw, report_data)

        wb.save(dest)
        return dest

    @classmethod
    def _style_headers(cls, ws, row_idx: int, cols_count: int):
        for col in range(1, cols_count + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cls.THIN_BORDER

    @classmethod
    def _populate_summary(cls, ws, report: Dict[str, Any]):
        ws.cell(row=1, column=1, value=report.get("title", "Business Intelligence Report")).font = cls.TITLE_FONT
        ws.cell(row=2, column=1, value=f"Period: {report.get('period', '')} | Database: {report.get('database_name', '')} | Generated: {report.get('generated_at', '')}").font = cls.REGULAR_FONT

        ws.cell(row=4, column=1, value="Executive Observations").font = cls.BOLD_FONT
        row_cursor = 5
        for obs in report.get("observations", []):
            ws.cell(row=row_cursor, column=1, value=f"• {obs}").font = cls.REGULAR_FONT
            row_cursor += 1

        row_cursor += 1
        ws.cell(row=row_cursor, column=1, value="Database Overview").font = cls.BOLD_FONT
        row_cursor += 1
        catalog = report.get("catalog_summary", {})
        ws.cell(row=row_cursor, column=1, value=f"Total Tables: {catalog.get('table_count', 0)}").font = cls.REGULAR_FONT
        row_cursor += 1
        ws.cell(row=row_cursor, column=1, value=f"Total Records Profiled: {catalog.get('total_rows', 0):,}").font = cls.REGULAR_FONT
        ws.column_dimensions["A"].width = 80

    @classmethod
    def _populate_kpis(cls, ws, kpis: list):
        headers = ["KPI Name", "Formula", "Current Value", "Unit", "Previous Period", "MoM Growth %", "Target Value", "Target Achievement %"]
        for c_idx, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c_idx, value=h)
        cls._style_headers(ws, 1, len(headers))

        for r_idx, k in enumerate(kpis, start=2):
            ws.cell(row=r_idx, column=1, value=k.get("name", "")).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=2, value=k.get("formula", "")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=k.get("current_value", 0)).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=4, value=k.get("unit", "")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=5, value=k.get("previous_period_value", "-")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=f"{k.get('growth_rate_mom_pct', 0):+0.1f}%" if k.get("growth_rate_mom_pct") else "-").font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=7, value=k.get("target_value", "-")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=8, value=f"{k.get('target_achievement_pct', 0):.1f}%" if k.get("target_achievement_pct") else "-").font = cls.REGULAR_FONT

        for c in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            ws.column_dimensions[c].width = 22

    @classmethod
    def _populate_trends(cls, ws, trends: list):
        headers = ["Metric", "Trajectory", "Total Growth %", "Linear Slope", "R² Score", "Volatility (CV %)", "Statistical Explanation"]
        for c_idx, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c_idx, value=h)
        cls._style_headers(ws, 1, len(headers))

        for r_idx, t in enumerate(trends, start=2):
            ws.cell(row=r_idx, column=1, value=t.get("metric_name", "")).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=2, value=t.get("trend_direction", "")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=f"{t.get('total_growth_percentage', 0):+0.1f}%").font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=4, value=t.get("linear_slope", 0)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=5, value=t.get("r_squared", 0)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=t.get("volatility_coefficient_variation", 0)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=7, value=t.get("explanation", "")).font = cls.REGULAR_FONT

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["G"].width = 60

    @classmethod
    def _populate_anomalies(cls, ws, anomalies: list):
        headers = ["Entity", "Metric", "Severity", "Observed Value", "Normal Min", "Normal Max", "Deviation %", "Detection Method", "Audit Rationale"]
        for c_idx, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c_idx, value=h)
        cls._style_headers(ws, 1, len(headers))

        for r_idx, a in enumerate(anomalies, start=2):
            ws.cell(row=r_idx, column=1, value=str(a.get("entity", ""))).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=2, value=str(a.get("metric", ""))).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=str(a.get("severity", ""))).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=4, value=a.get("observed_value", 0)).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=5, value=a.get("normal_range_min", 0)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=a.get("normal_range_max", 0)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=7, value=str(a.get("deviation_pct", ""))).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=8, value=str(a.get("method", ""))).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=9, value=str(a.get("explanation", ""))).font = cls.REGULAR_FONT

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["I"].width = 60

    @classmethod
    def _populate_data_quality(cls, ws, dq: dict):
        ws.cell(row=1, column=1, value="Overall Quality Score").font = cls.BOLD_FONT
        ws.cell(row=1, column=2, value=f"{dq.get('overall_quality_score', 100)}/100 ({dq.get('overall_grade', 'GOOD')})").font = cls.BOLD_FONT

        headers = ["Table Name", "Table Score", "Grade", "Checks Evaluated", "Violations Count", "Scoring Rationale"]
        for c_idx, h in enumerate(headers, start=1):
            ws.cell(row=3, column=c_idx, value=h)
        cls._style_headers(ws, 3, len(headers))

        for r_idx, t in enumerate(dq.get("tables", []), start=4):
            ws.cell(row=r_idx, column=1, value=t.get("table_name", "")).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=2, value=t.get("quality_score", 100)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=t.get("quality_grade", "GOOD")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=4, value=t.get("total_checks", 0)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=5, value=t.get("violations_count", 0)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=t.get("scoring_explanation", "")).font = cls.REGULAR_FONT

        for c in ["A", "B", "C", "D", "E"]:
            ws.column_dimensions[c].width = 20
        ws.column_dimensions["F"].width = 60

    @classmethod
    def _populate_raw_results(cls, ws, report: dict):
        headers = ["Category", "Source Object", "Status / Severity", "Detail Statement / Value"]
        for c_idx, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c_idx, value=h)
        cls._style_headers(ws, 1, len(headers))

        r_idx = 2
        for rule in report.get("rule_violations", []):
            ws.cell(row=r_idx, column=1, value="Business Rule").font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=2, value=rule.get("rule_name", "")).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=3, value=rule.get("severity", "HIGH")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=4, value=f"{rule.get('alert_message', '')} ({rule.get('violation_count', 0)} occurrences)").font = cls.REGULAR_FONT
            r_idx += 1

        for ins in report.get("insights", []):
            ws.cell(row=r_idx, column=1, value="Diagnostic Insight").font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=2, value=ins.get("table_name", "")).font = cls.BOLD_FONT
            ws.cell(row=r_idx, column=3, value=ins.get("direction", "")).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=4, value=ins.get("headline", "")).font = cls.REGULAR_FONT
            r_idx += 1

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 70
